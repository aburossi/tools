import streamlit as st
import markdown
from docx import Document
from io import BytesIO

def markdown_to_docx(markdown_text):
    # Convert Markdown to HTML
    html_text = markdown.markdown(markdown_text)
    
    # Create a new Document
    doc = Document()
    
    # Add a heading and content to the document
    for line in html_text.splitlines():
        if line.startswith('<h1>'):
            doc.add_heading(line[4:-5], level=1)
        elif line.startswith('<h2>'):
            doc.add_heading(line[4:-5], level=2)
        elif line.startswith('<ul>'):
            continue
        elif line.startswith('<li>'):
            doc.add_paragraph(line[4:-5], style='ListBullet')
        elif line.startswith('<p>'):
            doc.add_paragraph(line[3:-4])
        elif line.startswith('<strong>'):
            doc.add_paragraph(line[8:-9], style='Bold')
        elif line.startswith('<em>'):
            doc.add_paragraph(line[4:-5], style='Italic')
    
    # Save the document to a BytesIO object
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Streamlit app
st.title("Markdown to DOCX Converter")

st.write("Enter your Markdown text below:")

# Text area for markdown input
markdown_text = st.text_area("Markdown Input", height=300)

if st.button("Convert to DOCX"):
    # Convert and save as DOCX
    docx_buffer = markdown_to_docx(markdown_text)
    
    # Display download link for the DOCX
    st.download_button(
        label="Download DOCX",
        data=docx_buffer,
        file_name="output.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
